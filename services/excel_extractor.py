"""
Excel file extraction - extracts tables from .xlsx and .xls files.

Mirrors the interface of services/extractor.py for PDFs.
"""

import io
import re
from typing import Any

import pandas as pd


def extract_as_tables(excel_bytes: bytes, filename: str = "") -> list[dict]:
    """Extract tables from an Excel file.

    Args:
        excel_bytes: Raw bytes of the Excel file
        filename: Original filename (used to detect .xls vs .xlsx)

    Returns:
        List of table dicts in the same format as PDF extractor:
        {type, title, subtitle, headers/data/rows, page}

        For Excel, 'page' represents the sheet index (1-based).
        Cell values may be strings or dicts with 'value' and 'formula' keys.
    """
    is_xls = filename.lower().endswith('.xls') and not filename.lower().endswith('.xlsx')

    if is_xls:
        return _extract_xls(excel_bytes)
    else:
        return _extract_xlsx(excel_bytes)


def _fix_xlsx_stylesheet(excel_bytes: bytes) -> bytes | None:
    """Patch invalid aRGB color values in xl/styles.xml so openpyxl can load the file.

    Returns patched bytes if any fixes were applied, or None if no bad colors
    were found (so the caller can skip an unnecessary retry).

    Some tools write 6-char (RGB) or truncated hex colors instead of the 8-char
    aRGB format openpyxl requires.  We normalise them in-place inside the ZIP.
    """
    import zipfile

    _STYLES_PATH = "xl/styles.xml"
    # Match rgb="..." attributes that are NOT exactly 8 hex chars
    _BAD_COLOR = re.compile(r'(rgb=")([0-9A-Fa-f]{1,7}|[0-9A-Fa-f]{9,})(")')

    def _fix_color(m: re.Match) -> str:
        val = m.group(2).upper()
        if len(val) == 6:       # plain RGB without alpha → prepend FF (full opacity)
            val = "FF" + val
        elif len(val) < 8:      # truncated → right-pad with F (preserves intent)
            val = val + "F" * (8 - len(val))
        else:                   # too long → truncate to 8
            val = val[:8]
        return m.group(1) + val + m.group(3)

    buf_in = io.BytesIO(excel_bytes)
    buf_out = io.BytesIO()
    fixed_any = False
    with zipfile.ZipFile(buf_in, "r") as zin, zipfile.ZipFile(buf_out, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            raw = zin.read(item.filename)
            if item.filename == _STYLES_PATH:
                text = raw.decode("utf-8", errors="replace")
                patched_text, n = _BAD_COLOR.subn(_fix_color, text)
                if n:
                    fixed_any = True
                    raw = patched_text.encode("utf-8")
            zout.writestr(item, raw)

    return buf_out.getvalue() if fixed_any else None


def _extract_xlsx(excel_bytes: bytes) -> list[dict]:
    """Extract tables from .xlsx files using openpyxl."""
    import openpyxl

    def _load(data: bytes):
        return openpyxl.load_workbook(io.BytesIO(data), data_only=False)

    try:
        wb = _load(excel_bytes)
    except ValueError as e:
        if "password" in str(e).lower() or "encrypted" in str(e).lower():
            raise ValueError("File is password-protected. Please upload an unprotected version.")
        # Stylesheet XML issue — attempt in-memory patch and retry once
        try:
            patched = _fix_xlsx_stylesheet(excel_bytes)
        except Exception:
            patched = None
        if patched is None:
            raise ValueError(f"Unable to read Excel file. Please verify it's a valid .xlsx file. Error: {e}")
        try:
            wb = _load(patched)
        except Exception as e2:
            raise ValueError(f"Unable to read Excel file. Please verify it's a valid .xlsx file. Error: {e2}")
    except Exception as e:
        raise ValueError(f"Unable to read Excel file. Please verify it's a valid .xlsx file. Error: {e}")

    all_tables = []

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]
        tables = _extract_tables_from_sheet(ws, sheet_idx, sheet_name)
        all_tables.extend(tables)

    wb.close()
    return all_tables


def _extract_xls(excel_bytes: bytes) -> list[dict]:
    """Extract tables from .xls files using xlrd."""
    try:
        import xlrd
    except ImportError:
        raise ImportError("xlrd is required for .xls files. Install with: pip install xlrd")

    try:
        wb = xlrd.open_workbook(file_contents=excel_bytes)
    except xlrd.biffh.XLRDError as e:
        if "password" in str(e).lower() or "encrypted" in str(e).lower():
            raise ValueError("File is password-protected. Please upload an unprotected version.")
        raise ValueError(f"Unable to read Excel file. Please verify it's a valid .xls file. Error: {e}")

    all_tables = []

    for sheet_idx in range(wb.nsheets):
        sheet = wb.sheet_by_index(sheet_idx)
        sheet_name = sheet.name
        tables = _extract_tables_from_xls_sheet(sheet, sheet_idx + 1, sheet_name)
        all_tables.extend(tables)

    return all_tables


def _extract_tables_from_sheet(ws, sheet_idx: int, sheet_name: str) -> list[dict]:
    """Extract tables from an openpyxl worksheet by detecting table boundaries."""
    from services.table_parser import TableParser

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    if max_row == 0 or max_col == 0:
        return []

    # Build a grid of cell data with formulas
    grid = []
    for row_idx in range(1, max_row + 1):
        row_data = []
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell_data = _get_cell_data(cell)
            row_data.append(cell_data)
        grid.append(row_data)

    # Detect table regions separated by blank rows
    table_regions = _detect_table_regions(grid)

    tables = []
    for region in table_regions:
        df = _region_to_dataframe(region)
        if df is not None and not df.empty:
            parsed = TableParser.parse(df)
            if parsed:
                # Prefix title with sheet name for disambiguation
                original_title = parsed.get("title", "")
                if original_title:
                    parsed["title"] = f"{sheet_name}: {original_title}"
                else:
                    parsed["title"] = sheet_name
                parsed["page"] = sheet_idx

                # Preserve formula information in the parsed structure
                _add_formula_info(parsed, region)
                tables.append(parsed)

    return tables


def _extract_tables_from_xls_sheet(sheet, sheet_idx: int, sheet_name: str) -> list[dict]:
    """Extract tables from an xlrd sheet."""
    from services.table_parser import TableParser

    max_row = sheet.nrows
    max_col = sheet.ncols

    if max_row == 0 or max_col == 0:
        return []

    # Build a grid of cell data (xlrd doesn't preserve formulas)
    grid = []
    for row_idx in range(max_row):
        row_data = []
        for col_idx in range(max_col):
            try:
                value = sheet.cell_value(row_idx, col_idx)
                if isinstance(value, float) and value == int(value):
                    value = int(value)
                row_data.append(str(value) if value != "" else "")
            except Exception:
                row_data.append("")
        grid.append(row_data)

    # Detect table regions
    table_regions = _detect_table_regions(grid)

    tables = []
    for region in table_regions:
        df = _region_to_dataframe(region)
        if df is not None and not df.empty:
            parsed = TableParser.parse(df)
            if parsed:
                original_title = parsed.get("title", "")
                if original_title:
                    parsed["title"] = f"{sheet_name}: {original_title}"
                else:
                    parsed["title"] = sheet_name
                parsed["page"] = sheet_idx
                tables.append(parsed)

    return tables


def _get_cell_data(cell) -> str | dict:
    """Extract cell value and formula from an openpyxl cell.

    Returns:
        - String for plain values
        - Dict with 'value' and 'formula' keys for formula cells
    """
    raw_value = cell.value
    formula = None

    # Check for formula first - data_type == 'f' indicates a formula
    if hasattr(cell, 'data_type') and cell.data_type == 'f':
        # When data_only=False, the value IS the formula string
        if isinstance(raw_value, str) and raw_value.startswith('='):
            formula = raw_value
            # Try to get cached value if available
            # Unfortunately with data_only=False we don't have the calculated value
            # We'll use the formula string as a display value indicator
            value = f"[Formula: {formula}]"
        else:
            value = str(raw_value) if raw_value is not None else ""
    elif raw_value is None:
        value = ""
    elif isinstance(raw_value, float):
        # Convert floats that are whole numbers to int for cleaner display
        if raw_value == int(raw_value):
            value = str(int(raw_value))
        else:
            value = str(raw_value)
    else:
        value = str(raw_value)

    if formula:
        return {"value": value, "formula": formula}
    return value


def _is_row_empty(row: list) -> bool:
    """Check if a row is empty (all cells are empty strings)."""
    for cell in row:
        if isinstance(cell, dict):
            if cell.get("value", "").strip():
                return False
        elif isinstance(cell, str) and cell.strip():
            return False
    return True


def _detect_table_regions(grid: list[list]) -> list[list[list]]:
    """Detect separate table regions in the grid based on blank rows.

    A table boundary is detected when there are 2+ consecutive blank rows.
    """
    if not grid:
        return []

    regions = []
    current_region = []
    blank_count = 0

    for row in grid:
        if _is_row_empty(row):
            blank_count += 1
            if blank_count >= 2 and current_region:
                # End of current table region
                regions.append(current_region)
                current_region = []
        else:
            if blank_count == 1 and current_region:
                # Single blank row - might be part of the table
                current_region.append([""] * len(row))
            blank_count = 0
            current_region.append(row)

    # Don't forget the last region
    if current_region:
        regions.append(current_region)

    return regions


def _region_to_dataframe(region: list[list]) -> pd.DataFrame | None:
    """Convert a table region to a pandas DataFrame.

    Extracts just the string values for DataFrame construction.
    Formula information is preserved separately.
    """
    if not region:
        return None

    # Extract string values from cells (handle dict cells)
    rows = []
    for row in region:
        row_values = []
        for cell in row:
            if isinstance(cell, dict):
                row_values.append(cell.get("value", ""))
            else:
                row_values.append(cell)
        rows.append(row_values)

    if not rows:
        return None

    # Normalize column count
    max_cols = max(len(r) for r in rows)
    normalized = [r + [""] * (max_cols - len(r)) for r in rows]

    return pd.DataFrame(normalized)


def _add_formula_info(parsed: dict, region: list[list]) -> None:
    """Add formula information to the parsed table structure.

    For columnar tables, adds formula info to row values.
    For key_value tables, adds formula info to data values.
    For matrix tables, adds formula info to row cell values.
    """
    # Build a map of (row, col) -> formula
    formula_map = {}
    for r_idx, row in enumerate(region):
        for c_idx, cell in enumerate(row):
            if isinstance(cell, dict) and cell.get("formula"):
                formula_map[(r_idx, c_idx)] = cell["formula"]

    if not formula_map:
        return  # No formulas to add

    t_type = parsed.get("type")

    if t_type == "columnar":
        # For columnar tables, the rows are dicts keyed by header
        # We need to find which original cells map to which header values
        headers = parsed.get("headers", [])
        rows = parsed.get("rows", [])

        # Assuming first row (after title rows) is headers
        # and subsequent rows are data
        # This is a simplified approach - formulas in data cells
        for row_dict in rows:
            for header in headers:
                val = row_dict.get(header, "")
                if isinstance(val, str):
                    # Check if any formula exists for this value
                    for (r_idx, c_idx), formula in formula_map.items():
                        cell = region[r_idx][c_idx]
                        cell_val = cell.get("value", "") if isinstance(cell, dict) else cell
                        if cell_val == val:
                            row_dict[header] = {"value": val, "formula": formula}
                            break

    elif t_type == "key_value":
        data = parsed.get("data", {})
        for key, val in list(data.items()):
            if isinstance(val, str):
                for (r_idx, c_idx), formula in formula_map.items():
                    cell = region[r_idx][c_idx]
                    cell_val = cell.get("value", "") if isinstance(cell, dict) else cell
                    if cell_val == val:
                        data[key] = {"value": val, "formula": formula}
                        break

    elif t_type == "matrix":
        rows = parsed.get("rows", {})
        for row_label, row_data in rows.items():
            for col, val in list(row_data.items()):
                if isinstance(val, str):
                    for (r_idx, c_idx), formula in formula_map.items():
                        cell = region[r_idx][c_idx]
                        cell_val = cell.get("value", "") if isinstance(cell, dict) else cell
                        if cell_val == val:
                            row_data[col] = {"value": val, "formula": formula}
                            break


def _handle_merged_cells(ws) -> dict:
    """Get a map of merged cell ranges for unmerging.

    Returns a dict mapping each cell in a merged range to the merged value.
    """
    merged_map = {}

    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        # Get the value from the top-left cell
        value = ws.cell(row=min_row, column=min_col).value

        # Map all cells in the range to this value
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                merged_map[(row, col)] = value

    return merged_map
