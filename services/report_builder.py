import io
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.comparison import ComparisonResult
from services.utils import strip_filename_metadata

_FILLS = {
    "match":            PatternFill("solid", fgColor="FFC6EFCE"),  # Green
    "mismatch":         PatternFill("solid", fgColor="FFFFC7CE"),  # Red
    "formula_mismatch": PatternFill("solid", fgColor="FFFFA500"),  # Orange
    "only_in_a":        PatternFill("solid", fgColor="FFFFEB9C"),  # Yellow
    "only_in_b":        PatternFill("solid", fgColor="FFFFEB9C"),  # Yellow
    "missing":          PatternFill("solid", fgColor="FFD9D9D9"),  # Grey
}
_HEADER_FILL  = PatternFill("solid", fgColor="FF4472C4")
_SECTION_FILL = PatternFill("solid", fgColor="FF7030A0")   # purple for table section headers
_HEADER_FONT  = Font(bold=True, color="FFFFFFFF")
_SECTION_FONT = Font(bold=True, color="FFFFFFFF")
_N_COLS = 4   # Field | Doc A | Doc B | Status


def _extract_report_id(name_a: str, name_b: str) -> str:
    for name in (name_b, name_a):
        cleaned = strip_filename_metadata(name)
        if cleaned:
            return cleaned
    return "Report"


def _sheet_name(title: str, used: set[str]) -> str:
    safe = re.sub(r'[\\/*?:\[\]]', '-', title)[:31]
    if safe not in used:
        used.add(safe)
        return safe
    base = safe[:27]
    for n in range(2, 100):
        candidate = f"{base} ({n})"[:31]
        if candidate not in used:
            used.add(candidate)
            return candidate
    return safe


def _write_col_header(ws, row: int, name_a: str, name_b: str) -> None:
    cols = ["Field", name_a, name_b, "Status"]
    for col_idx, name in enumerate(cols, 1):
        cell = ws.cell(row=row, column=col_idx, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _write_section_header(ws, row: int, label: str) -> None:
    cell = ws.cell(row=row, column=1, value=label)
    cell.fill = _SECTION_FILL
    cell.font = _SECTION_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=_N_COLS)


def _write_diff_row(ws, row_idx: int, d, name_a: str, name_b: str) -> None:
    field_label = d.field.split(" › ", 1)[1] if " › " in d.field else d.field

    # Build value strings - include formula info for formula_mismatch
    val_a_str = d.value_a or ""
    val_b_str = d.value_b or ""

    if d.status == "formula_mismatch":
        if d.formula_a:
            val_a_str = f"{val_a_str} [{d.formula_a}]"
        if d.formula_b:
            val_b_str = f"{val_b_str} [{d.formula_b}]"
        status_label = "formula differs"
    elif d.status == "only_in_a":
        status_label = f"missing in {name_b}"
    elif d.status == "only_in_b":
        status_label = f"missing in {name_a}"
    else:
        status_label = d.status

    fill = _FILLS.get(d.status, _FILLS["mismatch"])
    for col_idx, val in enumerate([field_label, val_a_str, val_b_str, status_label], 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = fill


def _display_title(title: str) -> str:
    """Strip the '(p. N)' disambiguation suffix added internally for uniqueness."""
    return re.sub(r'\s*\(p\.\s*[\d?]+\)\s*$', '', title)


def _write_pair_sheet(ws, result: ComparisonResult, name_a: str, name_b: str) -> None:
    diffs_by_table: dict[str, list] = {}
    for d in result.diffs:
        diffs_by_table.setdefault(d.table, []).append(d)
    missing_by_title = {mt["title"]: mt for mt in result.missing_tables}

    # Freeze header row (col headers written at row 1 before the first section)
    _write_col_header(ws, row=1, name_a=name_a, name_b=name_b)
    ws.freeze_panes = ws.cell(row=2, column=1)

    current_row = 2
    for tname in result.table_order:
        if tname in missing_by_title:
            mt = missing_by_title[tname]
            page_str = f"  ·  p. {mt['page']}" if mt.get("page") else ""
            subtitle = result.table_subtitles.get(tname, "")
            subtitle_str = f"  —  {subtitle}" if subtitle else ""
            doc_missing = name_a if mt["missing_from"] == "A" else name_b
            label = f"{_display_title(tname)}{subtitle_str}{page_str}  —  MISSING IN {doc_missing}"
            _write_section_header(ws, current_row, label)
            current_row += 1
        else:
            table_diffs = diffs_by_table.get(tname, [])
            page = table_diffs[0].page if table_diffs else 0
            page_str = f"  ·  p. {page}" if page else ""
            subtitle = result.table_subtitles.get(tname, "")
            subtitle_str = f"  —  {subtitle}" if subtitle else ""
            failures = sum(1 for d in table_diffs if d.status != "match")
            status_str = "  —  OK" if failures == 0 else f"  —  {failures} issue(s)"
            label = f"{_display_title(tname)}{subtitle_str}{page_str}{status_str}"
            _write_section_header(ws, current_row, label)
            current_row += 1
            for d in table_diffs:
                _write_diff_row(ws, current_row, d, name_a, name_b)
                current_row += 1

        # blank separator between tables
        current_row += 1

    # Auto-fit columns
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col if not getattr(cell, 'merged', False)), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def build_excel_report(pairs: list[tuple[ComparisonResult, str, str]]) -> bytes:
    """Build a workbook with one sheet per comparison pair, named by report ID."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()

    for result, name_a, name_b in pairs:
        report_id = _extract_report_id(name_a, name_b)
        sname = _sheet_name(report_id, used_names)
        ws = wb.create_sheet(sname)
        _write_pair_sheet(ws, result, name_a, name_b)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
