"""
Generate Excel fixture files for testing Excel comparison.

Mirrors the data patterns from generate_fixture_pdfs.py.
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = Path(__file__).parent

# Styling
HEADER_FILL = PatternFill("solid", fgColor="1a3c5e")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALT_FILL = PatternFill("solid", fgColor="eaf1f8")
BORDER = Border(
    left=Side(style='thin', color='b0c4d8'),
    right=Side(style='thin', color='b0c4d8'),
    top=Side(style='thin', color='b0c4d8'),
    bottom=Side(style='thin', color='b0c4d8'),
)


def _style_header(ws, row: int, cols: int):
    """Apply header styling to a row."""
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER


def _style_data_row(ws, row: int, cols: int, alt: bool = False):
    """Apply data row styling."""
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        if alt:
            cell.fill = ALT_FILL
        cell.border = BORDER


# Shared table data (same as PDF fixtures)
TABLE_TITLE = "Channel Performance Metrics"
HEADER = ["Channel", "Region", "Viewers (M)", "Avg Watch (min)", "Peak Hour", "Ad Fill Rate"]

ROWS = [
    ["NBC East",         "Northeast",    "4.21",  "38",  "20:00", "94 %"],
    ["NBC West",         "West Coast",   "3.87",  "41",  "21:00", "92 %"],
    ["MSNBC",            "National",     "2.15",  "55",  "19:00", "88 %"],
    ["CNBC",             "National",     "1.93",  "48",  "09:00", "91 %"],
    ["USA Network",      "National",     "1.74",  "32",  "21:00", "89 %"],
    ["Bravo",            "National",     "1.41",  "29",  "21:00", "87 %"],
    ["E! Entertainment", "National",     "1.22",  "26",  "22:00", "85 %"],
    ["Oxygen",           "National",     "0.98",  "44",  "20:00", "83 %"],
    ["Syfy",             "National",     "0.87",  "61",  "21:00", "86 %"],
    ["NBC Sports",       "National",     "3.10",  "72",  "15:00", "96 %"],
    ["Golf Channel",     "National",     "0.45",  "90",  "13:00", "90 %"],
    ["Peacock Free",     "Digital",      "5.60",  "24",  "20:00", "78 %"],
    ["Peacock Premium",  "Digital",      "3.30",  "52",  "21:00", "95 %"],
    ["Telemundo East",   "Northeast",    "1.88",  "33",  "20:00", "82 %"],
    ["Telemundo West",   "West Coast",   "1.65",  "35",  "21:00", "81 %"],
]


def _write_table(ws, start_row: int, title: str, headers: list, rows: list) -> int:
    """Write a table to a worksheet. Returns the next available row."""
    current_row = start_row

    # Title
    ws.cell(row=current_row, column=1, value=title)
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=14)
    current_row += 1

    # Headers
    for col, header in enumerate(headers, 1):
        ws.cell(row=current_row, column=col, value=header)
    _style_header(ws, current_row, len(headers))
    current_row += 1

    # Data rows
    for i, row_data in enumerate(rows):
        for col, value in enumerate(row_data, 1):
            ws.cell(row=current_row, column=col, value=value)
        _style_data_row(ws, current_row, len(row_data), alt=(i % 2 == 1))
        current_row += 1

    return current_row + 1  # blank row after table


def build_simple_before():
    """Create simple_before.xlsx - basic single-sheet Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Performance"

    _write_table(ws, 1, TABLE_TITLE, HEADER, ROWS[:10])

    wb.save(OUT / "simple_before.xlsx")
    print("  simple_before.xlsx")


def build_simple_after():
    """Create simple_after.xlsx - same structure with some value differences."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Performance"

    # Modify some values
    modified_rows = [list(r) for r in ROWS[:10]]
    modified_rows[0][2] = "4.25"  # Slightly different viewer count
    modified_rows[3][5] = "93 %"  # Different ad fill rate

    _write_table(ws, 1, TABLE_TITLE, HEADER, modified_rows)

    wb.save(OUT / "simple_after.xlsx")
    print("  simple_after.xlsx")


def build_multi_sheet_before():
    """Create multi_sheet_before.xlsx - multiple sheets."""
    wb = openpyxl.Workbook()

    # Sheet 1: Performance Metrics
    ws1 = wb.active
    ws1.title = "Performance"
    _write_table(ws1, 1, TABLE_TITLE, HEADER, ROWS[:8])

    # Sheet 2: Financial Summary
    ws2 = wb.create_sheet("Financial")
    financial_headers = ["Category", "Q1", "Q2", "Q3", "Q4", "Total"]
    financial_rows = [
        ["Revenue", "1,200,000", "1,350,000", "1,400,000", "1,500,000", "5,450,000"],
        ["Expenses", "800,000", "850,000", "900,000", "950,000", "3,500,000"],
        ["Profit", "400,000", "500,000", "500,000", "550,000", "1,950,000"],
    ]
    _write_table(ws2, 1, "Quarterly Financial Summary", financial_headers, financial_rows)

    # Sheet 3: Key Value pairs
    ws3 = wb.create_sheet("Metadata")
    ws3.cell(row=1, column=1, value="Document Info")
    ws3.cell(row=1, column=1).font = Font(bold=True, size=14)
    kv_data = [
        ["Report Date", "2025-01-15"],
        ["Prepared By", "Analytics Team"],
        ["Version", "1.0"],
        ["Status", "Draft"],
    ]
    for i, (key, val) in enumerate(kv_data, 2):
        ws3.cell(row=i, column=1, value=key)
        ws3.cell(row=i, column=2, value=val)

    wb.save(OUT / "multi_sheet_before.xlsx")
    print("  multi_sheet_before.xlsx")


def build_multi_sheet_after():
    """Create multi_sheet_after.xlsx - same structure with differences."""
    wb = openpyxl.Workbook()

    # Sheet 1: Performance Metrics (with differences)
    ws1 = wb.active
    ws1.title = "Performance"
    modified_rows = [list(r) for r in ROWS[:8]]
    modified_rows[1][2] = "3.90"  # Different viewer count
    _write_table(ws1, 1, TABLE_TITLE, HEADER, modified_rows)

    # Sheet 2: Financial Summary (with differences)
    ws2 = wb.create_sheet("Financial")
    financial_headers = ["Category", "Q1", "Q2", "Q3", "Q4", "Total"]
    financial_rows = [
        ["Revenue", "1,200,000", "1,380,000", "1,400,000", "1,500,000", "5,480,000"],  # Q2 and Total differ
        ["Expenses", "800,000", "850,000", "900,000", "950,000", "3,500,000"],
        ["Profit", "400,000", "530,000", "500,000", "550,000", "1,980,000"],  # Q2 and Total differ
    ]
    _write_table(ws2, 1, "Quarterly Financial Summary", financial_headers, financial_rows)

    # Sheet 3: Key Value pairs (with differences)
    ws3 = wb.create_sheet("Metadata")
    ws3.cell(row=1, column=1, value="Document Info")
    ws3.cell(row=1, column=1).font = Font(bold=True, size=14)
    kv_data = [
        ["Report Date", "2025-01-20"],  # Different date
        ["Prepared By", "Analytics Team"],
        ["Version", "1.1"],  # Different version
        ["Status", "Final"],  # Different status
    ]
    for i, (key, val) in enumerate(kv_data, 2):
        ws3.cell(row=i, column=1, value=key)
        ws3.cell(row=i, column=2, value=val)

    wb.save(OUT / "multi_sheet_after.xlsx")
    print("  multi_sheet_after.xlsx")


def build_formula_diff_a():
    """Create formula_diff_a.xlsx - with formulas for totals."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"

    ws.cell(row=1, column=1, value="Budget Summary")
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)

    headers = ["Item", "Q1", "Q2", "Q3", "Q4", "Annual"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=header)
    _style_header(ws, 2, len(headers))

    # Data with formulas
    data = [
        ["Marketing", 100000, 120000, 110000, 130000],
        ["Engineering", 200000, 210000, 220000, 230000],
        ["Operations", 80000, 85000, 90000, 95000],
    ]

    for row_idx, row_data in enumerate(data, 3):
        ws.cell(row=row_idx, column=1, value=row_data[0])
        for col_idx, val in enumerate(row_data[1:], 2):
            ws.cell(row=row_idx, column=col_idx, value=val)
        # Formula for Annual total - using SUM
        ws.cell(row=row_idx, column=6, value=f"=SUM(B{row_idx}:E{row_idx})")
        _style_data_row(ws, row_idx, 6, alt=(row_idx % 2 == 0))

    # Total row with formulas
    total_row = len(data) + 3
    ws.cell(row=total_row, column=1, value="Total")
    for col in range(2, 7):
        col_letter = openpyxl.utils.get_column_letter(col)
        ws.cell(row=total_row, column=col, value=f"=SUM({col_letter}3:{col_letter}{total_row - 1})")
    _style_data_row(ws, total_row, 6)

    wb.save(OUT / "formula_diff_a.xlsx")
    print("  formula_diff_a.xlsx")


def build_formula_diff_b():
    """Create formula_diff_b.xlsx - same values but different formulas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"

    ws.cell(row=1, column=1, value="Budget Summary")
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)

    headers = ["Item", "Q1", "Q2", "Q3", "Q4", "Annual"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=header)
    _style_header(ws, 2, len(headers))

    # Same data with different formulas (individual adds instead of SUM)
    data = [
        ["Marketing", 100000, 120000, 110000, 130000],
        ["Engineering", 200000, 210000, 220000, 230000],
        ["Operations", 80000, 85000, 90000, 95000],
    ]

    for row_idx, row_data in enumerate(data, 3):
        ws.cell(row=row_idx, column=1, value=row_data[0])
        for col_idx, val in enumerate(row_data[1:], 2):
            ws.cell(row=row_idx, column=col_idx, value=val)
        # Different formula for Annual total - using addition
        ws.cell(row=row_idx, column=6, value=f"=B{row_idx}+C{row_idx}+D{row_idx}+E{row_idx}")
        _style_data_row(ws, row_idx, 6, alt=(row_idx % 2 == 0))

    # Total row with different formula structure
    total_row = len(data) + 3
    ws.cell(row=total_row, column=1, value="Total")
    for col in range(2, 7):
        col_letter = openpyxl.utils.get_column_letter(col)
        # Using addition instead of SUM
        ws.cell(row=total_row, column=col, value=f"={col_letter}3+{col_letter}4+{col_letter}5")
    _style_data_row(ws, total_row, 6)

    wb.save(OUT / "formula_diff_b.xlsx")
    print("  formula_diff_b.xlsx")


def build_empty_sheet():
    """Create a file with an empty sheet to test edge case."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Data"
    _write_table(ws1, 1, "Sample Data", ["A", "B", "C"], [["1", "2", "3"]])

    ws2 = wb.create_sheet("Empty")
    # Leave this sheet empty

    ws3 = wb.create_sheet("More Data")
    _write_table(ws3, 1, "More Sample Data", ["X", "Y"], [["10", "20"]])

    wb.save(OUT / "with_empty_sheet.xlsx")
    print("  with_empty_sheet.xlsx")


def build_merged_cells():
    """Create a file with merged cells to test edge case."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"

    # Title spanning multiple columns
    ws.cell(row=1, column=1, value="Quarterly Report 2025")
    ws.merge_cells("A1:D1")
    ws.cell(row=1, column=1).font = Font(bold=True, size=16)

    # Headers
    for col, header in enumerate(["Region", "Q1", "Q2", "Total"], 1):
        ws.cell(row=3, column=col, value=header)
    _style_header(ws, 3, 4)

    # Data
    data = [
        ["North", "100", "120", "220"],
        ["South", "80", "90", "170"],
        ["East", "110", "130", "240"],
    ]
    for row_idx, row_data in enumerate(data, 4):
        for col_idx, val in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
        _style_data_row(ws, row_idx, 4, alt=(row_idx % 2 == 0))

    wb.save(OUT / "merged_cells.xlsx")
    print("  merged_cells.xlsx")


if __name__ == "__main__":
    print(f"Writing Excel fixtures to {OUT}/")
    build_simple_before()
    build_simple_after()
    build_multi_sheet_before()
    build_multi_sheet_after()
    build_formula_diff_a()
    build_formula_diff_b()
    build_empty_sheet()
    build_merged_cells()
    print("Done.")
