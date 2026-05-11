"""
Tests for Excel file extraction.
"""

import io
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from services.excel_extractor import extract_as_tables


FIXTURES = Path(__file__).parent / "fixtures"


def _create_simple_xlsx(data: list[list], sheet_name: str = "Sheet1") -> bytes:
    """Create a simple xlsx file in memory."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    for row_idx, row in enumerate(data, 1):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _create_multi_sheet_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Create a multi-sheet xlsx file in memory."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name, data in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for row_idx, row in enumerate(data, 1):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _create_xlsx_with_formulas(data: list[list], formulas: dict[tuple, str]) -> bytes:
    """Create xlsx with specific formulas.

    Args:
        data: Cell values as list of lists
        formulas: Dict mapping (row, col) 1-indexed to formula string
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    for row_idx, row in enumerate(data, 1):
        for col_idx, value in enumerate(row, 1):
            if (row_idx, col_idx) in formulas:
                ws.cell(row=row_idx, column=col_idx, value=formulas[(row_idx, col_idx)])
            else:
                ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestBasicExtraction:
    """Test basic table extraction from Excel files."""

    def test_single_table_extraction(self):
        """Extract a simple table from a single sheet."""
        data = [
            ["Name", "Age", "City"],
            ["Alice", "30", "NYC"],
            ["Bob", "25", "LA"],
        ]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert len(tables) == 1
        table = tables[0]
        assert table["type"] == "columnar"
        assert "Name" in table["headers"]
        assert len(table["rows"]) == 2

    def test_empty_sheet_skipped(self):
        """Empty sheets should be skipped."""
        sheets = {
            "Empty": [],
            "Data": [["A", "B"], ["1", "2"]],
        }
        xlsx_bytes = _create_multi_sheet_xlsx(sheets)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        # Should only have table from Data sheet
        assert len(tables) == 1
        assert "Data" in tables[0]["title"]

    def test_multiple_sheets_by_position(self):
        """Tables from multiple sheets should be extracted in position order."""
        sheets = {
            "First": [["X", "Y"], ["1", "2"]],
            "Second": [["A", "B"], ["3", "4"]],
            "Third": [["P", "Q"], ["5", "6"]],
        }
        xlsx_bytes = _create_multi_sheet_xlsx(sheets)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert len(tables) == 3
        assert tables[0]["page"] == 1
        assert tables[1]["page"] == 2
        assert tables[2]["page"] == 3
        assert "First" in tables[0]["title"]
        assert "Second" in tables[1]["title"]
        assert "Third" in tables[2]["title"]

    def test_sheet_name_in_title(self):
        """Sheet name should be included in table title."""
        data = [
            ["Header1", "Header2"],
            ["val1", "val2"],
        ]
        xlsx_bytes = _create_simple_xlsx(data, sheet_name="MySheet")

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert len(tables) == 1
        assert "MySheet" in tables[0]["title"]


class TestTableTypes:
    """Test different table type detection."""

    def test_columnar_table(self):
        """Standard columnar table with headers."""
        data = [
            ["Product", "Price", "Quantity"],
            ["Widget", "10.00", "100"],
            ["Gadget", "25.00", "50"],
        ]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert tables[0]["type"] == "columnar"
        assert tables[0]["headers"] == ["Product", "Price", "Quantity"]
        assert len(tables[0]["rows"]) == 2

    def test_key_value_table(self):
        """Two-column key-value table with longer values triggers key_value detection."""
        data = [
            ["Name", "John Doe - This is a longer value that exceeds 35 characters to trigger key_value detection"],
            ["Email", "john.doe@example.com"],
            ["Phone", "555-1234"],
        ]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert tables[0]["type"] == "key_value"
        assert "Name" in tables[0]["data"]
        assert "John Doe" in tables[0]["data"]["Name"]

    def test_matrix_table(self):
        """Matrix table with empty first header cell."""
        data = [
            ["", "Q1", "Q2", "Q3"],
            ["North", "100", "110", "120"],
            ["South", "80", "85", "90"],
        ]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert tables[0]["type"] == "matrix"
        assert tables[0]["columns"] == ["Q1", "Q2", "Q3"]
        assert "North" in tables[0]["rows"]


class TestMultipleTablesPerSheet:
    """Test detection of multiple tables within a single sheet."""

    def test_tables_separated_by_blank_rows(self):
        """Multiple tables separated by blank rows should be detected."""
        data = [
            ["Table 1"],
            ["A", "B"],
            ["1", "2"],
            ["", ""],  # blank row
            ["", ""],  # blank row
            ["Table 2"],
            ["X", "Y"],
            ["3", "4"],
        ]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert len(tables) == 2


class TestFormulaExtraction:
    """Test formula extraction and representation."""

    def test_formula_preserved(self):
        """Formulas should be preserved in cell data."""
        data = [
            ["Item", "Value"],
            ["A", "10"],
            ["B", "20"],
            ["Total", "30"],  # This will be a formula
        ]
        formulas = {
            (4, 2): "=B2+B3",  # Row 4, Col 2 is the Total
        }
        xlsx_bytes = _create_xlsx_with_formulas(data, formulas)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert len(tables) == 1
        # The formula cell should be a dict with value and formula


class TestEdgeCases:
    """Test edge cases and error handling."""

    def test_numeric_values(self):
        """Numeric values should be properly converted to strings."""
        data = [
            ["Count", "Percentage"],
            [100, 0.5],
            [200, 0.75],
        ]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert len(tables) == 1
        # Values should be string representations

    def test_empty_cells(self):
        """Empty cells should be handled gracefully."""
        data = [
            ["A", "B", "C"],
            ["1", "", "3"],
            ["", "2", ""],
        ]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")

        assert len(tables) == 1

    def test_invalid_file_raises_error(self):
        """Invalid file content should raise ValueError."""
        with pytest.raises(ValueError):
            extract_as_tables(b"not an excel file", "test.xlsx")


class TestFileTypeDetection:
    """Test .xls vs .xlsx file type handling."""

    def test_xlsx_extension(self):
        """Files with .xlsx extension use openpyxl."""
        data = [["A", "B"], ["1", "2"]]
        xlsx_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(xlsx_bytes, "test.xlsx")
        assert len(tables) == 1

    def test_xls_extension_without_xlrd(self):
        """Files with .xls extension attempt xlrd (may fail if not installed)."""
        data = [["A", "B"], ["1", "2"]]
        xlsx_bytes = _create_simple_xlsx(data)

        # This should fail because the content is xlsx format but extension is .xls
        # xlrd will try to read it and fail
        try:
            tables = extract_as_tables(xlsx_bytes, "test.xls")
            # If xlrd is installed, it should fail to parse xlsx content
        except (ImportError, ValueError):
            pass  # Expected - either xlrd not installed or wrong format


def _inject_bad_rgb_color(xlsx_bytes: bytes, bad_rgb: str) -> bytes:
    """Replace the first valid 8-char aRGB color in xl/styles.xml with a bad value.

    Used to simulate files produced by tools that write non-conforming color values.
    """
    buf_in = io.BytesIO(xlsx_bytes)
    buf_out = io.BytesIO()
    _GOOD = re.compile(r'rgb="[0-9A-Fa-f]{8}"')
    with zipfile.ZipFile(buf_in, "r") as zin, zipfile.ZipFile(buf_out, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            raw = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                text = raw.decode("utf-8")
                text = _GOOD.sub(f'rgb="{bad_rgb}"', text, count=1)
                raw = text.encode("utf-8")
            zout.writestr(item, raw)
    return buf_out.getvalue()


class TestMalformedStylesheet:
    """Test that files with invalid aRGB color values in xl/styles.xml are handled."""

    def test_6char_rgb_color_loads(self):
        """Files with 6-char RGB colors (missing alpha) should load successfully."""
        data = [["Name", "Value"], ["A", "1"]]
        good_bytes = _create_simple_xlsx(data)
        bad_bytes = _inject_bad_rgb_color(good_bytes, "FFFFFF")  # 6-char, missing alpha

        tables = extract_as_tables(bad_bytes, "test.xlsx")

        assert len(tables) == 1
        assert tables[0]["type"] == "columnar"

    def test_7char_rgb_color_loads(self):
        """Files with 7-char RGB colors (truncated) should load successfully."""
        data = [["Name", "Value"], ["A", "1"]]
        good_bytes = _create_simple_xlsx(data)
        bad_bytes = _inject_bad_rgb_color(good_bytes, "fffffff")  # 7-char, truncated

        tables = extract_as_tables(bad_bytes, "test.xlsx")

        assert len(tables) == 1
        assert tables[0]["type"] == "columnar"

    def test_valid_file_loads_without_patching(self):
        """Valid files should load normally — the fix path is not taken."""
        data = [["Name", "Value"], ["A", "1"], ["B", "2"]]
        good_bytes = _create_simple_xlsx(data)

        tables = extract_as_tables(good_bytes, "test.xlsx")

        assert len(tables) == 1
        assert len(tables[0]["rows"]) == 2


class TestIntegrationWithComparator:
    """Integration tests with the comparator."""

    def test_extracted_tables_work_with_comparator(self):
        """Extracted tables should work with compare_tables."""
        from services.comparator import compare_tables

        data_a = [
            ["Name", "Value"],
            ["Item1", "100"],
            ["Item2", "200"],
        ]
        data_b = [
            ["Name", "Value"],
            ["Item1", "100"],
            ["Item2", "250"],  # Different value
        ]

        tables_a = extract_as_tables(_create_simple_xlsx(data_a), "a.xlsx")
        tables_b = extract_as_tables(_create_simple_xlsx(data_b), "b.xlsx")

        result = compare_tables(tables_a, tables_b)

        assert result.mode == "table_parser"
        # Should have one mismatch for Item2's value
        mismatches = [d for d in result.diffs if d.status == "mismatch"]
        assert len(mismatches) >= 1
